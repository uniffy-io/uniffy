"""Office document text extraction (DOCX, XLSX, PPTX)."""

import io

from uniffy.core.extraction.types import ExtractionResult


def extract_docx_text(data: bytes, max_chars: int) -> ExtractionResult:
    """Extract DOCX paragraphs and table cell contents."""
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts: list[str] = []
    total_chars = 0
    truncated = False

    for para in doc.paragraphs:
        text = para.text.strip()
        if text:
            parts.append(text)
            total_chars += len(text)
            if total_chars >= max_chars:
                truncated = True
                break

    if not truncated:
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    row_text = "\t".join(cells)
                    parts.append(row_text)
                    total_chars += len(row_text)
                    if total_chars >= max_chars:
                        truncated = True
                        break
            if truncated:
                break

    text = "\n".join(parts)
    if len(text) > max_chars:
        text = text[:max_chars]
        truncated = True

    word_count = len(text.split()) if text.strip() else 0

    return ExtractionResult(
        text=text,
        word_count=word_count,
        truncated=truncated,
    )


def extract_xlsx_text(data: bytes, max_chars: int) -> ExtractionResult:
    """Extract XLSX cell values per sheet, tab-separated, with sheet name headers."""
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    total_chars = 0
    truncated = False
    page_count = len(wb.sheetnames)

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            parts.append(f"--- Sheet: {sheet_name} ---")
            total_chars += len(sheet_name) + 15

            for row in ws.iter_rows(values_only=True):
                cells = [str(cell) if cell is not None else "" for cell in row]
                if any(c for c in cells):
                    row_text = "\t".join(cells)
                    parts.append(row_text)
                    total_chars += len(row_text)
                    if total_chars >= max_chars:
                        truncated = True
                        break
            if truncated:
                break
    finally:
        wb.close()

    text = "\n".join(parts)
    if len(text) > max_chars:
        text = text[:max_chars]
        truncated = True

    word_count = len(text.split()) if text.strip() else 0

    return ExtractionResult(
        text=text,
        page_count=page_count,
        word_count=word_count,
        truncated=truncated,
    )


def extract_pptx_text(data: bytes, max_chars: int) -> ExtractionResult:
    """Extract PPTX slide text shapes and speaker notes, with slide number headers."""
    from pptx import Presentation

    prs = Presentation(io.BytesIO(data))
    parts: list[str] = []
    total_chars = 0
    truncated = False

    for slide_num, slide in enumerate(prs.slides, 1):
        parts.append(f"--- Slide {slide_num} ---")
        total_chars += 15

        for shape in slide.shapes:
            if shape.has_text_frame:
                for paragraph in shape.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        parts.append(text)
                        total_chars += len(text)
                        if total_chars >= max_chars:
                            truncated = True
                            break
            if truncated:
                break

        if not truncated and slide.has_notes_slide:
            notes_frame = slide.notes_slide.notes_text_frame
            if notes_frame:
                notes_text = notes_frame.text.strip()
                if notes_text:
                    parts.append(f"[Notes: {notes_text}]")
                    total_chars += len(notes_text)

        if truncated:
            break

    text = "\n".join(parts)
    if len(text) > max_chars:
        text = text[:max_chars]
        truncated = True

    word_count = len(text.split()) if text.strip() else 0

    return ExtractionResult(
        text=text,
        page_count=len(prs.slides),
        word_count=word_count,
        truncated=truncated,
    )
