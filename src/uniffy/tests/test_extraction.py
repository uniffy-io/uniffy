"""Tests for the shared extraction module."""

import io

import pytest

from uniffy.core.extraction import (
    EXTRACTABLE_MIME_TYPES,
    ExtractionResult,
    UnsupportedFormatError,
    can_extract,
    extract_text,
)
from uniffy.core.extraction.markup import extract_html_text, extract_rtf_text
from uniffy.core.extraction.text import extract_csv_text, extract_plain_text


class TestCanExtract:
    """Tests for the can_extract function."""

    def test_pdf_supported(self) -> None:
        assert can_extract("application/pdf") is True

    def test_docx_supported(self) -> None:
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert can_extract(mime) is True

    def test_xlsx_supported(self) -> None:
        mime = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        assert can_extract(mime) is True

    def test_pptx_supported(self) -> None:
        mime = "application/vnd.openxmlformats-officedocument.presentationml.presentation"
        assert can_extract(mime) is True

    def test_html_supported(self) -> None:
        assert can_extract("text/html") is True

    def test_plain_text_supported(self) -> None:
        assert can_extract("text/plain") is True

    def test_json_supported(self) -> None:
        assert can_extract("application/json") is True

    def test_csv_supported(self) -> None:
        assert can_extract("text/csv") is True

    def test_text_wildcard_fallback(self) -> None:
        assert can_extract("text/x-custom-type") is True

    def test_image_not_supported(self) -> None:
        assert can_extract("image/png") is False

    def test_video_not_supported(self) -> None:
        assert can_extract("video/mp4") is False

    def test_binary_not_supported(self) -> None:
        assert can_extract("application/octet-stream") is False

    def test_mime_with_parameters(self) -> None:
        assert can_extract("text/plain; charset=utf-8") is True

    def test_case_insensitive(self) -> None:
        assert can_extract("Application/PDF") is True


class TestExtractTextRouting:
    """Tests for MIME routing in extract_text."""

    def test_plain_text(self) -> None:
        data = b"Hello, world!"
        result = extract_text(data, "text/plain")
        assert isinstance(result, ExtractionResult)
        assert result.text == "Hello, world!"

    def test_text_fallback_for_unknown_text_subtype(self) -> None:
        data = b"some content"
        result = extract_text(data, "text/x-unknown")
        assert result.text == "some content"

    def test_unsupported_format_raises(self) -> None:
        with pytest.raises(UnsupportedFormatError):
            extract_text(b"\x00\x01\x02", "image/png")

    def test_json_extracted_as_text(self) -> None:
        data = b'{"key": "value"}'
        result = extract_text(data, "application/json")
        assert '"key"' in result.text

    def test_mime_with_parameters(self) -> None:
        data = b"content"
        result = extract_text(data, "text/plain; charset=utf-8")
        assert result.text == "content"


class TestPlainTextExtraction:
    """Tests for plain text extraction."""

    def test_utf8_text(self) -> None:
        data = b"Hello World"
        result = extract_plain_text(data, 1000)
        assert result.text == "Hello World"
        assert result.word_count == 2

    def test_bom_handling(self) -> None:
        data = b"\xef\xbb\xbfHello BOM"
        result = extract_plain_text(data, 1000)
        assert result.text == "Hello BOM"
        assert not result.text.startswith("\ufeff")

    def test_max_chars_truncation(self) -> None:
        data = ("a" * 500).encode("utf-8")
        result = extract_plain_text(data, 100)
        assert len(result.text) == 100
        assert result.truncated is True

    def test_invalid_utf8_replaced(self) -> None:
        data = b"Hello \xff World"
        result = extract_plain_text(data, 1000)
        assert "Hello" in result.text
        assert "World" in result.text

    def test_empty_text(self) -> None:
        result = extract_plain_text(b"", 1000)
        assert result.text == ""
        assert result.word_count == 0


class TestCsvExtraction:
    """Tests for CSV/TSV extraction."""

    def test_basic_csv(self) -> None:
        data = b"name,age,city\nAlice,30,NYC\nBob,25,LA"
        result = extract_csv_text(data, 10000)
        assert "Alice" in result.text
        assert "Bob" in result.text
        assert "\t" in result.text

    def test_csv_max_chars(self) -> None:
        rows = [f"row{i},value{i}" for i in range(1000)]
        data = "\n".join(rows).encode("utf-8")
        result = extract_csv_text(data, 100)
        assert len(result.text) <= 100
        assert result.truncated is True

    def test_csv_with_bom(self) -> None:
        data = b"\xef\xbb\xbfname,value\na,1"
        result = extract_csv_text(data, 10000)
        assert "name" in result.text


class TestHtmlExtraction:
    """Tests for HTML text extraction."""

    def test_basic_html(self) -> None:
        data = b"<html><body><p>Hello World</p></body></html>"
        result = extract_html_text(data, 10000)
        assert "Hello World" in result.text

    def test_strips_tags(self) -> None:
        data = b"<h1>Title</h1><p>Paragraph</p>"
        result = extract_html_text(data, 10000)
        assert "<h1>" not in result.text
        assert "Title" in result.text
        assert "Paragraph" in result.text

    def test_strips_script(self) -> None:
        data = b"<p>Text</p><script>alert('xss')</script><p>More</p>"
        result = extract_html_text(data, 10000)
        assert "alert" not in result.text
        assert "Text" in result.text
        assert "More" in result.text

    def test_strips_style(self) -> None:
        data = b"<style>.body{color:red}</style><p>Content</p>"
        result = extract_html_text(data, 10000)
        assert "color" not in result.text
        assert "Content" in result.text

    def test_max_chars(self) -> None:
        data = ("<p>" + "x" * 500 + "</p>").encode()
        result = extract_html_text(data, 100)
        assert len(result.text) == 100
        assert result.truncated is True


class TestRtfExtraction:
    """Tests for RTF text extraction."""

    def test_basic_rtf(self) -> None:
        data = rb"{\rtf1\ansi Hello World}"
        result = extract_rtf_text(data, 10000)
        assert "Hello World" in result.text

    def test_strips_control_words(self) -> None:
        data = rb"{\rtf1\ansi\deff0 Some \b bold \b0 text}"
        result = extract_rtf_text(data, 10000)
        assert "Some" in result.text
        assert "text" in result.text
        assert "\\b" not in result.text


class TestPdfExtraction:
    """Tests for PDF text extraction."""

    def test_basic_pdf(self) -> None:
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "Hello PDF World")
        pdf_bytes = doc.tobytes()
        doc.close()

        result = extract_text(pdf_bytes, "application/pdf")
        assert "Hello PDF World" in result.text
        assert result.page_count == 1

    def test_multipage_pdf(self) -> None:
        import fitz

        doc = fitz.open()
        for i in range(3):
            page = doc.new_page()
            page.insert_text((72, 72), f"Page {i + 1} content")
        pdf_bytes = doc.tobytes()
        doc.close()

        result = extract_text(pdf_bytes, "application/pdf")
        assert result.page_count == 3
        assert "Page 1" in result.text
        assert "Page 3" in result.text

    def test_pdf_max_chars(self) -> None:
        import fitz

        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "x" * 500)
        pdf_bytes = doc.tobytes()
        doc.close()

        result = extract_text(pdf_bytes, "application/pdf", max_chars=100)
        assert len(result.text) <= 100
        assert result.truncated is True


class TestDocxExtraction:
    """Tests for DOCX text extraction."""

    def test_basic_docx(self) -> None:
        from docx import Document

        doc = Document()
        doc.add_paragraph("Hello DOCX World")
        doc.add_paragraph("Second paragraph")
        buf = io.BytesIO()
        doc.save(buf)

        result = extract_text(
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert "Hello DOCX World" in result.text
        assert "Second paragraph" in result.text

    def test_docx_with_table(self) -> None:
        from docx import Document

        doc = Document()
        doc.add_paragraph("Before table")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "A1"
        table.cell(0, 1).text = "B1"
        table.cell(1, 0).text = "A2"
        table.cell(1, 1).text = "B2"
        buf = io.BytesIO()
        doc.save(buf)

        result = extract_text(
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        assert "Before table" in result.text
        assert "A1" in result.text
        assert "B2" in result.text


class TestXlsxExtraction:
    """Tests for XLSX text extraction."""

    def test_basic_xlsx(self) -> None:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws["A1"] = "Header1"
        ws["B1"] = "Header2"
        ws["A2"] = "Value1"
        ws["B2"] = 42
        buf = io.BytesIO()
        wb.save(buf)

        result = extract_text(
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        assert "Header1" in result.text
        assert "42" in result.text

    def test_xlsx_multiple_sheets(self) -> None:
        from openpyxl import Workbook

        wb = Workbook()
        ws1 = wb.active
        ws1.title = "Sales"
        ws1["A1"] = "Revenue"
        ws2 = wb.create_sheet("Costs")
        ws2["A1"] = "Expenses"
        buf = io.BytesIO()
        wb.save(buf)

        result = extract_text(
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        assert "Sales" in result.text
        assert "Revenue" in result.text
        assert "Costs" in result.text
        assert "Expenses" in result.text
        assert result.page_count == 2


class TestPptxExtraction:
    """Tests for PPTX text extraction."""

    def test_basic_pptx(self) -> None:
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        title = slide.shapes.title
        title.text = "Slide Title"
        body = slide.placeholders[1]
        body.text = "Slide body content"
        buf = io.BytesIO()
        prs.save(buf)

        result = extract_text(
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        )
        assert "Slide Title" in result.text
        assert "Slide body content" in result.text
        assert result.page_count == 1


class TestExtractableMimeTypes:
    """Tests for the EXTRACTABLE_MIME_TYPES frozenset."""

    def test_contains_pdf(self) -> None:
        assert "application/pdf" in EXTRACTABLE_MIME_TYPES

    def test_contains_docx(self) -> None:
        mime = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        assert mime in EXTRACTABLE_MIME_TYPES

    def test_does_not_contain_images(self) -> None:
        assert "image/png" not in EXTRACTABLE_MIME_TYPES
        assert "image/jpeg" not in EXTRACTABLE_MIME_TYPES

    def test_is_frozenset(self) -> None:
        assert isinstance(EXTRACTABLE_MIME_TYPES, frozenset)
